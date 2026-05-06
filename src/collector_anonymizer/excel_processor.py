"""Excel Processor: de-anonymize AWS Transform Excel exports.

Scans all cells across all worksheets in an Excel workbook, replacing
anonymized values (including those embedded within surrounding text) with
their original values from the mapping store. Preserves formatting, formulas,
sheet names, and workbook structure.
"""

from pathlib import Path

from openpyxl import load_workbook

from collector_anonymizer.mapping_store import MappingStore


def deanonymize_excel(
    input_path: Path, output_path: Path, mapping: MappingStore
) -> int:
    """Scan all cells across all worksheets and replace anonymized values.

    For each cell that contains a string value, every anonymized value found
    in the reverse mapping is replaced with its original. Replacements work
    both for standalone cell values and for anonymized values embedded within
    surrounding text (e.g. ``"Migrate server-0001 to AWS"`` becomes
    ``"Migrate PROD-WEB-01 to AWS"``).

    Parameters
    ----------
    input_path:
        Path to the input ``.xlsx`` file.
    output_path:
        Path where the de-anonymized workbook will be saved.
    mapping:
        A :class:`MappingStore` whose reverse mapping is used for lookups.

    Returns
    -------
    int
        The total number of individual replacements made across all cells.
    """
    wb = load_workbook(str(input_path))

    # Build a flat lookup: anonymized -> original, sorted longest-first so
    # that longer anonymized tokens are matched before shorter substrings.
    reverse_flat: dict[str, str] = {}
    for category_map in mapping.reverse.values():
        reverse_flat.update(category_map)

    sorted_keys = sorted(reverse_flat.keys(), key=len, reverse=True)

    replacement_count = 0

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue

                original_text = cell.value
                new_text = original_text

                for anon_value in sorted_keys:
                    if anon_value in new_text:
                        count = new_text.count(anon_value)
                        new_text = new_text.replace(
                            anon_value, reverse_flat[anon_value]
                        )
                        replacement_count += count

                if new_text != original_text:
                    cell.value = new_text

    wb.save(str(output_path))
    return replacement_count
